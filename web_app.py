import os
import re
import pdfplumber
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from flask import Flask, request, render_template, send_file, jsonify
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = os.getenv('UPLOAD_FOLDER', 'uploads')
app.config['MAX_CONTENT_LENGTH'] = int(os.getenv('MAX_CONTENT_LENGTH', 16 * 1024 * 1024))  # 16MB max file size
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'your-secret-key-here')

# Ensure upload folder exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

def process_pdf(pdf_path, bank_type):
    transactions = []
    try:
        print(f"Processing PDF: {pdf_path} for bank type: {bank_type}")
        with pdfplumber.open(pdf_path) as pdf:
            print(f"PDF opened successfully. Number of pages: {len(pdf.pages)}")
            if bank_type == 'CMI':
                transactions = process_cmi_pdf(pdf)
            elif bank_type == 'BMCE':
                print("Processing BMCE PDF...")
                transactions = process_bmce_pdf(pdf)
                print(f"BMCE processing complete. Found {len(transactions)} transactions")
            elif bank_type == 'BP':
                transactions = process_bp_pdf(pdf)
            elif bank_type == 'BARID':
                transactions = process_barid_pdf(pdf)
            else:
                raise Exception(f"Unsupported bank type: {bank_type}")
        
        return transactions
    except Exception as e:
        print(f"Error in process_pdf: {str(e)}")
        raise Exception(f"Error processing PDF: {str(e)}")

def process_cmi_pdf(pdf):
    transactions = []
    for page in pdf.pages:
        text = page.extract_text()
        if not text:
            continue
        
        current_block = ""
        for line in text.split('\n'):
            line = line.strip()
            
            if line.startswith("ACHAT REMISE"):
                if current_block:
                    result = process_block(current_block)
                    if result:
                        transactions.append(result)
                current_block = line + "\n"
            elif current_block:
                current_block += line + "\n"
        
        if current_block:
            result = process_block(current_block)
            if result:
                transactions.append(result)
    
    return transactions

def process_bmce_pdf(pdf):
    transactions = []
    account_info = {}
    
    print("Starting BMCE PDF processing...")
    
    def clean_amount(amount_str):
        if not amount_str:
            return 0.0
        # Remove spaces and convert to float
        cleaned = amount_str.replace(' ', '').replace(',', '.')
        try:
            return float(cleaned)
        except ValueError:
            print(f"Error converting amount: {amount_str}")
            return 0.0
    
    def extract_amount(text):
        # Look for amount pattern: digits (possibly with spaces) followed by comma and 2 digits
        amount_match = re.search(r'((?:\d{1,3}\s)*\d+,\d{2})', text)
        if amount_match:
            return clean_amount(amount_match.group(1))
        return 0.0
    
    def is_debit_operation(description):
        debit_keywords = [
            'CASH POOLING DIRECT',
            'PRELEVEMENT',
            'RETRAIT',
            'COMMISSION',
            'TVA SUR',
            'FRAIS'
        ]
        return any(keyword in description.upper() for keyword in debit_keywords)
    
    def is_valid_date(day, month):
        try:
            # Check if month is valid (1-12)
            month_num = int(month)
            day_num = int(day)
            if not (1 <= month_num <= 12):
                return False
            # Check if day is valid for the given month
            datetime(2024, month_num, day_num)
            return True
        except ValueError:
            return False
    
    for page_num, page in enumerate(pdf.pages):
        text = page.extract_text()
        if not text:
            print(f"Page {page_num+1} has no text content")
            continue
        
        print(f"Processing page {page_num+1}")
        
        # Split text into lines and process each line
        lines = text.split('\n')
        current_transaction = None
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # Check if line starts with a date pattern (DD MM)
            date_match = re.match(r'^(\d{2})\s+(\d{2})\s+(.+)', line)
            if date_match:
                day, month, rest = date_match.groups()
                
                # Validate the date before processing
                if not is_valid_date(day, month):
                    # If not a valid date, treat as continuation if we have a current transaction
                    if current_transaction:
                        if not re.match(r'^\d+,\d{2}$', line):  # Don't append if line is just an amount
                            current_transaction['LIBELLE'] += " " + line
                    continue
                
                # If we have a pending transaction, save it
                if current_transaction:
                    transactions.append(current_transaction)
                
                # Start new transaction
                current_transaction = {
                    'DATE_OPERATION': f"{day}/{month}/2024",
                    'LIBELLE': rest.strip(),
                    'DATE_VALEUR': f"{day}/{month}/2024",  # Default to same date
                    'DEBIT': 0.0,
                    'CREDIT': 0.0
                }
            elif current_transaction:
                # Look for value date pattern (DD MM) in the continuation line
                val_date_match = re.search(r'(\d{2})\s+(\d{2})\s*$', line)
                if val_date_match:
                    val_day, val_month = val_date_match.groups()
                    if is_valid_date(val_day, val_month):
                        current_transaction['DATE_VALEUR'] = f"{val_day}/{val_month}/2024"
                        # Remove the date from the line before processing amount
                        line = line[:val_date_match.start()].strip()
                
                # Look for amount in the line
                amount = extract_amount(line)
                if amount > 0:
                    if is_debit_operation(current_transaction['LIBELLE']):
                        current_transaction['DEBIT'] = amount
                    else:
                        current_transaction['CREDIT'] = amount
                
                # Append the rest of the line to the description
                if not re.match(r'^\d+,\d{2}$', line):  # Don't append if line is just an amount
                    current_transaction['LIBELLE'] += " " + line
        
        # Don't forget to add the last transaction
        if current_transaction:
            transactions.append(current_transaction)
    
    # Clean up transaction descriptions
    for trans in transactions:
        # Remove multiple spaces and clean up the description
        trans['LIBELLE'] = ' '.join(trans['LIBELLE'].split())
        # Remove any trailing amounts from the description
        trans['LIBELLE'] = re.sub(r'\s+\d+,\d{2}\s*$', '', trans['LIBELLE'])
        # Remove any trailing dates
        trans['LIBELLE'] = re.sub(r'\s+\d{2}\s+\d{2}\s*$', '', trans['LIBELLE'])
    
    # Sort transactions by date
    transactions.sort(key=lambda x: datetime.strptime(x['DATE_OPERATION'], '%d/%m/%Y'))
    
    print(f"BMCE processing complete. Found {len(transactions)} transactions")
    return transactions

def process_bp_pdf(pdf):
    # Placeholder for Banque Populaire processing
    # This will be implemented when you provide the Banque Populaire PDF format
    raise Exception("Banque Populaire processing not yet implemented")

def process_barid_pdf(pdf):
    # Placeholder for Barid Bank processing
    # This will be implemented when you provide the Barid Bank PDF format
    raise Exception("Barid Bank processing not yet implemented")

def process_block(block):
    facture_match = re.search(r'ACHAT REMISE TPE N°\s*:\s*(\d+)', block)
    date_match = re.search(r'DU\s*:\s*(\d{2}/\d{2}/\d{2})', block)
    
    lines = block.split('\n')
    remise = commissions = tva = solde = None
    
    for i, line in enumerate(lines):
        if "TOTAL REMISE (DH)" in line:
            value = extract_value(line)
            if not value and i+1 < len(lines):
                value = extract_value(lines[i+1])
            remise = value
        elif "TOTAL COMMISSIONS HT" in line:
            value = extract_value(line)
            if not value and i+1 < len(lines):
                value = extract_value(lines[i+1])
            commissions = value
        elif "TOTAL TVA SUR COMMISSIONS" in line:
            value = extract_value(line)
            if not value and i+1 < len(lines):
                value = extract_value(lines[i+1])
            tva = value
        elif "SOLDE NET REMISE" in line:
            value = extract_value(line)
            if not value and i+1 < len(lines):
                value = extract_value(lines[i+1])
            solde = value
    
    if all([facture_match, date_match, remise, commissions, tva, solde]):
        try:
            return {
                'TPE': facture_match.group(1),
                'Date': date_match.group(1),
                'Total Remise (DH)': float(remise.replace(',', '')),
                'Total Commissions HT': float(commissions.replace(',', '')),
                'Total TVA': float(tva.replace(',', '')),
                'Solde Net Remise': float(solde.replace(',', ''))
            }
        except ValueError:
            return None
    return None

def extract_value(line):
    match = re.search(r'([\d,]+\.\d{2})\s*$', line.strip())
    return match.group(1) if match else None

def create_excel(transactions):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    
    # Check if we have BMCE transactions (they have different headers)
    if transactions and 'DATE_OPERATION' in transactions[0]:
        headers = ['DATE OPERATION', 'LIBELLE', 'DATE VALEUR', 'DEBIT', 'CREDIT']
        
        # Add account information if available
        if 'RIB' in transactions[0]:
            ws.append(['RELEVE D\'IDENTITE BANCAIRE', transactions[0]['RIB']])
        if 'PERIODE' in transactions[0]:
            ws.append(['PERIODE D\'EXTRAIT DE COMPTE', transactions[0]['PERIODE']])
        ws.append([])  # Empty row for separation
        
        # Add headers
        ws.append(headers)
        
        # Apply bold font to header cells
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.font = Font(bold=True)
        
        # Add transaction data
        for transaction in transactions:
            ws.append([
                transaction['DATE_OPERATION'],
                transaction['LIBELLE'],
                transaction['DATE_VALEUR'],
                transaction['DEBIT'],
                transaction['CREDIT']
            ])
    else:
        # Original CMI format
        headers = ['N° TPE', 'Date', 'TOTAL REMISE (DH)', 
                  'TOTAL COMMISSIONS HT', 'TOTAL TVA SUR COMMISSIONS', 'SOLDE NET REMISE']
        ws.append(headers)
        
        # Apply bold font to header cells
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
        
        for transaction in transactions:
            ws.append([
                transaction['TPE'],
                transaction['Date'],
                transaction['Total Remise (DH)'],
                transaction['Total Commissions HT'],
                transaction['Total TVA'],
                transaction['Solde Net Remise']
            ])
    
    return wb

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    
    file = request.files['file']
    bank_type = request.form.get('bank_type', 'CMI')  # Default to CMI if not specified
    
    print(f"Received file: {file.filename} for bank type: {bank_type}")
    
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400
    
    if not file.filename.endswith('.pdf'):
        return jsonify({'error': 'File must be a PDF'}), 400
    
    try:
        filename = secure_filename(file.filename)
        pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        print(f"Saving file to: {pdf_path}")
        file.save(pdf_path)
        
        transactions = process_pdf(pdf_path, bank_type)
        if not transactions:
            return jsonify({'error': 'No transactions found in the PDF'}), 400
        
        # Generate Excel file
        now = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_filename = f"transaction_{now}.xlsx"
        excel_path = os.path.join(app.config['UPLOAD_FOLDER'], excel_filename)
        
        wb = create_excel(transactions)
        wb.save(excel_path)
        
        # Clean up PDF file
        os.remove(pdf_path)
        
        return jsonify({
            'success': True,
            'message': f'Successfully processed {len(transactions)} transactions',
            'filename': excel_filename
        })
        
    except Exception as e:
        print(f"Error in upload_file: {str(e)}")
        return jsonify({'error': str(e)}), 500

@app.route('/download/<filename>')
def download_file(filename):
    try:
        return send_file(
            os.path.join(app.config['UPLOAD_FOLDER'], filename),
            as_attachment=True
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 404

if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    app.run(host='0.0.0.0', port=port) 